from pathlib import Path
from docx import Document

OUT = Path("datasets/docx_runtime/benign")
OUT.mkdir(parents=True, exist_ok=True)

documents = [
    ("report", "Monthly Business Report", "This document contains a normal business report."),
    ("meeting", "Meeting Notes", "These are ordinary meeting notes and action items."),
    ("invoice", "Invoice Summary", "This document contains a normal invoice summary."),
    ("letter", "Business Letter", "This is a normal business correspondence document."),
    ("resume", "Professional Resume", "This document contains ordinary professional information."),
    ("project", "Project Plan", "This document contains a normal project planning outline."),
    ("training", "Training Material", "This document contains ordinary training material."),
    ("minutes", "Meeting Minutes", "This document contains standard meeting minutes."),
    ("proposal", "Project Proposal", "This document contains a normal project proposal."),
    ("notes", "General Notes", "This document contains ordinary text and notes."),
]

for i in range(1, 6):
    for prefix, title, text in documents:
        path = OUT / f"{prefix}_{i:02d}.docx"

        doc = Document()
        doc.add_heading(title, level=1)
        doc.add_paragraph(text)
        doc.add_paragraph(
            f"Document identifier: {prefix}-{i:02d}. "
            "This file is part of a benign synthetic research dataset."
        )

        doc.save(path)

print(f"Created {len(list(OUT.glob('*.docx')))} benign DOCX files.")
print(f"Location: {OUT.resolve()}")